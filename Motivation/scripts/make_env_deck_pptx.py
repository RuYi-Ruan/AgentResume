"""Generate the environment-overview slide deck (docs/slides/environment_overview.pptx).

Mirrors docs/slides/environment_overview.html: diagrams over text, one visual per slide.
Run:  D:\\Anaconda3\\python.exe scripts/make_env_deck_pptx.py
"""

from __future__ import annotations

import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

OUT = pathlib.Path("docs/slides/environment_overview.pptx")

INK = RGBColor(0x2A, 0x23, 0x20)
INK_SOFT = RGBColor(0x6B, 0x5F, 0x58)
PAPER = RGBColor(0xF7, 0xF3, 0xEC)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
CARD_LINE = RGBColor(0xED, 0xE5, 0xDA)
TERRA = RGBColor(0xB8, 0x50, 0x42)
TERRA_D = RGBColor(0x8E, 0x3A, 0x2F)
TERRA_L = RGBColor(0xE8, 0xCF, 0xC8)
SAGE = RGBColor(0x7F, 0xA0, 0x8E)
SAGE_L = RGBColor(0xDC, 0xE8, 0xE1)
GOLD = RGBColor(0xE2, 0xB4, 0x57)
GOLD_L = RGBColor(0xF6, 0xE7, 0xC6)
DARK = RGBColor(0x24, 0x1F, 0x1C)
DARK_CARD = RGBColor(0x33, 0x2B, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ONION = RGBColor(0xE2, 0xC8, 0x8F)
POT = RGBColor(0xC9, 0xA0, 0xA0)
DISH = RGBColor(0x9F, 0xB8, 0xCF)
SERV = RGBColor(0xA7, 0xBE, 0xAE)
WALL = RGBColor(0x6B, 0x5F, 0x58)

FONT = "Microsoft YaHei"


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def box(slide, x, y, w, h, text="", size=18, color=INK, bold=False,
        fill=None, line=None, radius=True, align=PP_ALIGN.LEFT, font=FONT):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.12
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
    return shp


def circle(slide, x, y, d, text="", size=14, fill=TERRA, color=WHITE, bold=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = FONT
    return shp


def line(slide, x, y, w, color=RGBColor(0xC9, 0xB8, 0xAC), h=0.03, dash=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    if dash:
        shp.line.fill.solid()
    return shp


def text(slide, x, y, w, h, s, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.text = s
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
    return tb


def bg(slide, dark=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(12192000), Emu(6858000))
    shp.fill.solid()
    shp.fill.fore_color.rgb = DARK if dark else PAPER
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def header(slide, kicker, title, dark=False, accent=TERRA):
    text(slide, 0.72, 0.42, 6, 0.32, kicker, size=13, color=accent if not dark else RGBColor(0xD9, 0xCC, 0xC2), bold=True)
    text(slide, 0.72, 0.8, 11.9, 0.9, title, size=30, color=WHITE if dark else INK, bold=True)


def footer(slide, s, dark=False):
    text(slide, 0.72, 6.95, 11.0, 0.3, s, size=11, color=RGBColor(0xB9, 0xAA, 0xA0) if dark else INK_SOFT)


# ---------------------------------------------------------------- slides

def slide_title(prs):
    s = blank(prs)
    bg(s, dark=True)
    text(s, 0.72, 0.9, 11.9, 0.35, "AgentResume · 环境总览", size=13,
         color=RGBColor(0xD9, 0xCC, 0xC2), bold=True)
    text(s, 0.72, 1.5, 10.5, 1.9, "过时的印象，\n会让人读错伙伴吗？", size=44, color=WHITE, bold=True)
    text(s, 0.72, 3.5, 10.5, 0.5, "在 Overcooked 厨房里，让 Alice 变强，看 Bob 还能不能读懂她。",
         size=17, color=RGBColor(0xD9, 0xCC, 0xC2))
    circle(s, 1.9, 4.6, 1.3, "Alice", size=16, fill=TERRA)
    circle(s, 9.4, 4.6, 1.3, "Bob", size=16, fill=SAGE)
    line(s, 3.4, 5.15, 1.7)
    line(s, 7.6, 5.15, 1.7)
    text(s, 3.5, 4.68, 1.6, 0.3, "旧印象", size=12, color=RGBColor(0xB9, 0xAA, 0xA0))
    text(s, 7.7, 5.32, 1.7, 0.3, "真实动作", size=12, color=RGBColor(0xB9, 0xAA, 0xA0))
    text(s, 6.1, 4.95, 0.9, 0.6, "?", size=32, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    footer(s, "3 张地图 · 3 个训练种子 · 9 组实验 · 4050 次模型调用", dark=True)


def slide_board(prs):
    s = blank(prs)
    bg(s)
    header(s, "环境", "11 × 7 的厨房，两人一起出汤")
    cell = 0.42
    x0, y0 = 0.85, 1.95
    box(s, 0.72, 1.75, 6.1, 4.35, fill=CARD, line=CARD_LINE)
    walls = {(0, 2), (1, 2), (2, 2), (6, 2), (7, 2), (8, 2), (9, 2), (4, 4), (5, 4)}
    fac = {(1, 0): (ONION, "O"), (9, 6): (ONION, "O"),
           (3, 1): (POT, "P"), (5, 1): (POT, "P"), (7, 1): (POT, "P"),
           (2, 6): (DISH, "D"), (5, 6): (DISH, "D"), (8, 6): (DISH, "D"),
           (0, 3): (SERV, "S"), (0, 5): (SERV, "S"), (10, 3): (SERV, "S"), (10, 5): (SERV, "S")}
    for gx in range(12):
        line(s, x0 + gx * cell, y0, 0.008, color=RGBColor(0xE5, 0xDA, 0xCD), h=7 * cell, dash=False)
    for gy in range(8):
        line(s, x0, y0 + gy * cell, 11 * cell, color=RGBColor(0xE5, 0xDA, 0xCD), dash=False)
    for (wx, wy) in walls:
        box(s, x0 + wx * cell + 0.015, y0 + wy * cell + 0.015, cell - 0.03, cell - 0.03,
            fill=WALL, radius=False)
    for (fx, fy), (col, lab) in fac.items():
        box(s, x0 + fx * cell + 0.04, y0 + fy * cell + 0.04, cell - 0.08, cell - 0.08,
            lab, size=11, color=RGBColor(0x4A, 0x3B, 0x33), bold=True, fill=col, align=PP_ALIGN.CENTER)
    circle(s, x0 + 3 * cell + 0.06, y0 + 3 * cell + 0.06, 0.3, "A", size=11, fill=TERRA)
    circle(s, x0 + 6 * cell + 0.06, y0 + 5 * cell + 0.06, 0.3, "B", size=11, fill=SAGE)
    text(s, 0.85, 1.85, 5.0, 0.3, "每格 1 步 · 墙挡住视线与通路", size=12, color=RGBColor(0x8C, 0x7C, 0x72))

    steps = ["拿 3 个洋葱，放进锅里", "空手点火开煮", "煮 20 个 tick 出锅", "先拿盘子，再端汤", "送到出餐口，每碗 +20 分"]
    for i, st in enumerate(steps):
        y = 2.0 + i * 0.82
        box(s, 7.1, y, 5.5, 0.66, fill=CARD, line=CARD_LINE)
        circle(s, 7.22, y + 0.12, 0.42, str(i + 1), size=14, fill=TERRA)
        text(s, 7.85, y + 0.19, 4.5, 0.35, st, size=14, bold=True)
    text(s, 7.1, 6.2, 6.0, 0.3, "洋葱台 ×2 · 锅 ×3 · 盘子台 ×3 · 出餐口 ×4", size=12, color=INK_SOFT)
    footer(s, "三张地图 A / B / C 设施相同，只有墙不一样 · 一局 700 tick")


def slide_roles(prs):
    s = blank(prs)
    bg(s)
    header(s, "角色", "Alice 负责“变强”，Bob 负责“搞懂她”")
    box(s, 0.72, 1.75, 5.9, 3.6, fill=CARD, line=CARD_LINE)
    box(s, 6.9, 1.75, 5.7, 3.6, fill=CARD, line=CARD_LINE)
    circle(s, 1.0, 1.98, 0.62, "A", size=20, fill=TERRA)
    text(s, 1.78, 2.04, 4.6, 0.5, "Alice · 被研究的那个", size=20, bold=True)
    chain = [("182 维状态", TERRA_L, RGBColor(0x6B, 0x4A, 0x3F)),
             ("128", TERRA_L, RGBColor(0x6B, 0x4A, 0x3F)),
             ("128", TERRA_L, RGBColor(0x6B, 0x4A, 0x3F)),
             ("8 种意图", TERRA, WHITE)]
    cx = 1.05
    for (lab, col, fg) in chain:
        w = 1.35 if len(lab) < 6 else 1.6
        box(s, cx, 2.82, w, 0.6, lab, size=13, bold=True, color=fg, fill=col, align=PP_ALIGN.CENTER)
        cx += w + 0.18
        if lab != "8 种意图":
            text(s, cx - 0.13, 2.94, 0.3, 0.3, "→", size=15, color=RGBColor(0xB7, 0xA7, 0x9B), bold=True)
            cx += 0.24
    text(s, 1.05, 3.62, 5.3, 1.6,
         "输出 8 种高层意图：拿料 / 下锅 / 点火 / 拿盘 / 取汤 / 送餐 / 等锅 / 原地。\n"
         "走位和按键由“执行器”代劳——那是肌肉，不是决策。", size=13, color=INK_SOFT)
    circle(s, 7.18, 1.98, 0.62, "B", size=20, fill=SAGE)
    text(s, 7.96, 2.04, 4.4, 0.5, "Bob · 搭档兼读心者", size=20, bold=True)
    box(s, 7.25, 2.82, 1.7, 0.6, "供料", size=14, bold=True, color=RGBColor(0x6B, 0x4A, 0x3F),
        fill=RGBColor(0xEF, 0xE6, 0xDA), align=PP_ALIGN.CENTER)
    text(s, 9.02, 2.94, 0.4, 0.3, "/", size=16, color=RGBColor(0xB7, 0xA7, 0x9B), bold=True)
    box(s, 9.4, 2.82, 1.7, 0.6, "做服务", size=14, bold=True, color=RGBColor(0x6B, 0x4A, 0x3F),
        fill=RGBColor(0xEF, 0xE6, 0xDA), align=PP_ALIGN.CENTER)
    text(s, 10.62, 2.92, 0.4, 0.3, "+", size=16, color=TERRA, bold=True)
    box(s, 7.25, 3.55, 3.3, 0.55, "还要猜 Alice 想干嘛", size=14, bold=True, color=WHITE,
        fill=SAGE, align=PP_ALIGN.CENTER)
    text(s, 7.25, 4.2, 5.05, 1.05,
         "为什么只猜这两种？\n题目是有意挑出的岔路口：画面一样、第一步也一样——她要么守着锅（PRE），要么转身拿洋葱（FETCH）。其余 6 种意图的场面会露底。",
         size=11, color=RGBColor(0x7A, 0x53, 0x44))
    box(s, 0.72, 5.6, 11.9, 1.1, fill=RGBColor(0xFB, 0xED, 0xE9), line=TERRA)
    text(s, 0.98, 5.74, 11.4, 0.5,
         "关键：Alice 的真实意图不会喂给 Bob —— 他只能靠画面、动作和手里的印象。",
         size=19, bold=True, color=TERRA_D)
    text(s, 0.98, 6.24, 11.4, 0.35, "这三样东西的差别，就是实验唯一在变的量。",
         size=12, color=RGBColor(0x8C, 0x6A, 0x5E))


def slide_training(prs):
    s = blank(prs)
    bg(s)
    header(s, "Alice 的训练", "两个台阶：先模仿，再变强")
    box(s, 0.72, 1.9, 5.6, 2.3, fill=CARD, line=CARD_LINE)
    circle(s, 1.0, 2.15, 0.62, "1", size=20, fill=GOLD)
    text(s, 1.78, 2.2, 4.3, 0.4, "成长前 · 行为克隆", size=19, bold=True)
    text(s, 1.05, 3.0, 5.0, 0.9, "照着保守专家做：\n24 局训练 + 16 局独立测试，100 轮迭代。", size=13, color=INK_SOFT)
    box(s, 6.9, 1.9, 5.7, 2.3, fill=CARD, line=CARD_LINE)
    circle(s, 7.18, 2.15, 0.62, "2", size=20, fill=TERRA)
    text(s, 7.96, 2.2, 4.4, 0.4, "成长后 · 继续强化学习（PPO）", size=19, bold=True)
    text(s, 7.23, 3.0, 5.2, 0.9, "从成长前接着练：\n50 次更新 × 每次 8 局 × 700 tick。", size=13, color=INK_SOFT)
    # 阶梯
    line(s, 1.0, 5.6, 2.6, color=TERRA, h=0.05, dash=False)
    line(s, 3.6, 4.9, 0.05, color=TERRA, h=0.7, dash=False)
    line(s, 3.6, 4.9, 7.6, color=TERRA, h=0.05, dash=False)
    text(s, 3.9, 4.35, 7.3, 0.4, "能力门槛：平均多出 ≥ 1.0 碗，且 ≥ 17/32 个起点变好",
         size=14, color=TERRA_D, bold=True)
    text(s, 1.0, 5.75, 11.5, 0.4, "9 组全部过门槛：3 个训练种子 × 3 张地图。", size=14, color=INK_SOFT)


def slide_training_data(prs):
    s = blank(prs)
    bg(s)
    header(s, "训练 · 数据", "先跟专家抄，再自己练")
    nodes = [("规则专家跑 24 局", "每局 700 tick", TERRA),
             ("每个决策点记一条", "状态 → 意图", SAGE),
             ("3736 条样本", "每局约 156 条", GOLD),
             ("BC 100 轮", "照抄专家", TERRA)]
    x = 0.72
    for i, (title, sub, col) in enumerate(nodes):
        box(s, x, 1.6, 2.45, 1.4, fill=CARD, line=CARD_LINE)
        circle(s, x + 0.24, 1.72, 0.48, str(i + 1), size=14, fill=col)
        text(s, x + 0.24, 2.3, 2.05, 0.35, title, size=13, bold=True)
        text(s, x + 0.24, 2.6, 2.05, 0.3, sub, size=10, color=INK_SOFT)
        if i < 3:
            text(s, x + 2.5, 2.0, 0.4, 0.35, "→", size=18, color=RGBColor(0xC0, 0xAF, 0xA4), bold=True)
        x += 2.87
    box(s, 0.72, 3.25, 6.1, 2.9, fill=CARD, line=CARD_LINE)
    text(s, 0.98, 3.4, 5.6, 0.35, "训练数据的两条真样本", size=16, bold=True)
    box(s, 0.98, 3.9, 2.75, 1.15, fill=RGBColor(0xFB, 0xF7, 0xF1), line=RGBColor(0xE8, 0xDC, 0xCB))
    text(s, 1.12, 4.0, 2.5, 0.3, "标签 FETCH", size=13, bold=True, color=TERRA_D)
    text(s, 1.12, 4.3, 2.5, 0.65, "三口锅都空、两人空手、t=0\n→ 该去拿洋葱", size=10, color=INK_SOFT)
    box(s, 3.95, 3.9, 2.75, 1.15, fill=RGBColor(0xFB, 0xF7, 0xF1), line=RGBColor(0xE8, 0xDC, 0xCB))
    text(s, 4.09, 4.0, 2.5, 0.3, "标签 PRE", size=13, bold=True, color=TERRA_D)
    text(s, 4.09, 4.3, 2.5, 0.65, "1 号锅正在煮、空手、t=33\n→ 该在锅边等", size=10, color=INK_SOFT)
    text(s, 0.98, 5.25, 5.6, 0.8,
         "每条 = 182 维状态（位置 / 手持 / 三口锅 / 时间 / 已交付）+ 一个标签 + 一份“合法意图”掩码。",
         size=11, color=INK_SOFT)
    box(s, 7.05, 3.25, 5.55, 2.9, fill=CARD, line=CARD_LINE)
    text(s, 7.3, 3.4, 5.0, 0.35, "专家的标签都长什么样", size=16, bold=True)
    dist = [("PRE 等锅", 0.52, "52%", RGBColor(0xA9, 0x8D, 0x80)),
            ("HOLD 原地", 0.12, "12%", RGBColor(0xC6, 0xB3, 0xA8)),
            ("FETCH 拿料", 0.10, "10%", SAGE),
            ("PLACE 下锅", 0.10, "10%", RGBColor(0x9F, 0xB8, 0xCF)),
            ("其余四种", 0.16, "16%", RGBColor(0xC9, 0xBC, 0xB3))]
    dy = 3.8
    for label, frac, pct, col in dist:
        text(s, 7.3, dy + 0.04, 1.15, 0.3, label, size=11, bold=True,
             color=RGBColor(0x6B, 0x4A, 0x3F), align=PP_ALIGN.RIGHT)
        box(s, 8.55, dy, 3.85, 0.34, fill=RGBColor(0xEF, 0xE6, 0xDA))
        box(s, 8.55, dy, max(0.62, 3.85 * frac / 0.52), 0.34, pct, size=10, bold=True,
            color=WHITE, fill=col, align=PP_ALIGN.RIGHT)
        dy += 0.38
    text(s, 7.3, 5.74, 5.0, 0.4, "同一刻度：一半以上示范都是“在锅边等”，这是基础版爱站着的出处。",
         size=10, color=INK_SOFT)


def slide_training_why(prs):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    s = blank(prs)
    bg(s)
    header(s, "训练 · 变强", "变强不是被教出来的，是试出来又被留下的", accent=SAGE)
    box(s, 0.72, 1.8, 6.3, 3.5, fill=CARD, line=CARD_LINE)
    text(s, 0.98, 1.94, 5.8, 0.35, "验证集出汤：先掉下去，再爬上来", size=16, bold=True)
    chart_data = CategoryChartData()
    chart_data.categories = ["0", "5", "10", "15", "20", "25", "30", "35", "40", "45", "50"]
    chart_data.add_series("平均出汤", (14.0, 7.12, 15.0, 15.0, 14.38, 15.19, 11.31, 17.81, 18.56, 18.19, 20.0))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.95), Inches(2.35),
                            Inches(5.85), Inches(2.8), chart_data)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.size = Pt(10)
    plot = chart.plots[0]
    plot.has_data_labels = False
    series = plot.series[0]
    series.format.line.color.rgb = TERRA
    series.format.line.width = Pt(2.5)
    chart.value_axis.maximum_scale = 22.0
    chart.value_axis.minimum_scale = 5.0
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.size = Pt(9)
    text(s, 0.98, 5.05, 5.8, 0.3, "横轴：第几次更新（共 50 次）· 竖轴：16 个验证局平均出汤",
         size=10, color=INK_SOFT)
    box(s, 7.2, 1.8, 5.4, 3.5, fill=CARD, line=CARD_LINE)
    text(s, 7.45, 1.94, 5.0, 0.35, "行为被重新分配（同一批测试起点）", size=15, bold=True)
    bars = [("PRE 等锅", 648, 648, RGBColor(0xA9, 0x8D, 0x80)),
            ("", 56, 648, RGBColor(0xC9, 0xBC, 0xB3)),
            ("FETCH 拿料", 124, 648, SAGE),
            ("", 340, 648, TERRA)]
    by = 2.5
    for label, value, scale, col in bars:
        if label:
            text(s, 7.45, by + 0.04, 1.2, 0.3, label, size=11, bold=True, color=RGBColor(0x6B, 0x4A, 0x3F))
        box(s, 8.75, by, 3.6, 0.36, fill=RGBColor(0xEF, 0xE6, 0xDA))
        box(s, 8.75, by, max(0.5, 3.6 * value / scale), 0.36, str(value), size=10, bold=True,
            color=WHITE, fill=col, align=PP_ALIGN.RIGHT)
        by += 0.6
    text(s, 7.45, 4.9, 5.0, 0.4, "同一刻度（100% = 648 次）：等待几乎消失，拿料翻近三倍。",
         size=10, color=INK_SOFT)
    gaps = [("0.4%", "平时“拿料”的概率", RGBColor(0xEF, 0xE6, 0xDA), INK),
            ("16.9%", "探索时被抽中的概率", RGBColor(0xEF, 0xE6, 0xDA), INK),
            ("98.9%", "练成后的概率", RGBColor(0xFB, 0xED, 0xE9), TERRA_D)]
    for i, (big, sub, fill, col) in enumerate(gaps):
        left = 0.72 + i * 4.1
        box(s, left, 5.55, 3.85, 0.95, fill=fill, line=CARD_LINE)
        text(s, left + 0.2, 5.62, 3.5, 0.4, big, size=20, bold=True, color=col)
        text(s, left + 0.2, 6.08, 3.5, 0.3, sub, size=10, color=INK_SOFT)


def slide_group_matrix(prs):
    s = blank(prs)
    bg(s)
    header(s, "训练 · 覆盖", "9 组，9 个各自长大的 Alice")
    cells = [("地图 A", 1801, 14.0, 20.0, 6.0), ("地图 A", 1802, 14.0, 18.4, 4.4), ("地图 A", 1803, 13.6, 19.2, 5.6),
             ("地图 B", 1801, 14.5, 20.0, 5.5), ("地图 B", 1802, 14.4, 18.3, 3.9), ("地图 B", 1803, 14.6, 19.5, 4.9),
             ("地图 C", 1801, 15.7, 19.8, 4.1), ("地图 C", 1802, 15.4, 21.0, 5.6), ("地图 C", 1803, 15.1, 20.7, 5.6)]
    cw, ch = 3.9, 1.34
    for i, (mp, seed, pre, post, gain) in enumerate(cells):
        left = 0.72 + (i % 3) * (cw + 0.16)
        top = 1.62 + (i // 3) * (ch + 0.14)
        box(s, left, top, cw, ch, fill=CARD, line=CARD_LINE)
        text(s, left + 0.2, top + 0.1, 2.0, 0.28, mp, size=11, bold=True, color=RGBColor(0x8C, 0x7C, 0x72))
        text(s, left + 2.0, top + 0.1, 1.7, 0.28, f"种子 {seed}", size=11, bold=True,
             color=RGBColor(0x8C, 0x7C, 0x72), align=PP_ALIGN.RIGHT)
        text(s, left + 0.2, top + 0.42, 3.5, 0.5, f"{pre:.1f}  →  {post:.1f}", size=19, bold=True, color=INK)
        text(s, left + 0.2, top + 0.98, 3.5, 0.28, f"提升 +{gain:.1f}", size=11, bold=True,
             color=RGBColor(0x4E, 0x6B, 0x5B))
    notes = [("换地图", "厨房几何全变，门槛各自过"),
             ("换种子", "起点/示范/回合全换"),
             ("不共享权重", "post 接本组自己的 pre"),
             ("要的是复现", "9 组独立才叫结论稳")]
    x = 0.72
    for title, sub in notes:
        box(s, x, 6.0, 2.9, 0.95, fill=RGBColor(0xF2, 0xEB, 0xE1), line=CARD_LINE)
        text(s, x + 0.16, 6.08, 2.6, 0.3, title, size=12, bold=True)
        text(s, x + 0.16, 6.38, 2.6, 0.5, sub, size=9, color=INK_SOFT)
        x += 3.0
    text(s, 0.72, 7.05, 11.9, 0.3,
         "每组各自采集示范（3725–4273 条）、各自跑 400 局 PPO；9 组合计约 86 分钟。",
         size=10, color=INK_SOFT)


def slide_pipeline(prs):
    s = blank(prs)
    bg(s)
    header(s, "实验流程", "从训练到读数，一共六步", accent=SAGE)
    nodes = [("训练 Alice", "成长前 → 成长后", TERRA),
             ("攒印象", "Bob 旁观，只记看到的", SAGE),
             ("挑关键局面", "50 个 / 组，严格配对", GOLD),
             ("模型猜意图", "每题问 3 次，取多数", TERRA),
             ("续跑 200 tick", "看后面出几碗汤", SAGE),
             ("统计对比", "准确率为主，得分为辅", GOLD)]
    x = 0.72
    for i, (title, sub, col) in enumerate(nodes):
        box(s, x, 2.1, 1.78, 2.1, fill=CARD, line=CARD_LINE)
        circle(s, x + 0.58, 2.3, 0.62, str(i + 1), size=18, fill=col)
        text(s, x + 0.1, 3.05, 1.58, 0.4, title, size=15, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.1, 3.45, 1.58, 0.6, sub, size=11, color=INK_SOFT, align=PP_ALIGN.CENTER)
        x += 1.94
    line(s, 1.0, 5.15, 11.3, h=0.04)
    box(s, 0.72, 5.5, 11.85, 0.85, "唯一变量：Bob 手里那份印象（旧 / 新 / 无）——Alice 和局面都不动",
        size=15, bold=True, color=TERRA_D, fill=RGBColor(0xFB, 0xED, 0xE9))


def slide_key_moment(prs):
    s = blank(prs)
    bg(s)
    header(s, "关键局面", "同一幅画面，两种心思")
    for i, (cap, col, lab) in enumerate([
            ("成长前：先等着看锅", TERRA, "PRE"), ("成长后：转身去拿洋葱", SAGE, "FETCH")]):
        left = 0.72 + i * 6.2
        box(s, left, 1.9, 5.9, 3.6, fill=CARD, line=CARD_LINE)
        circle(s, left + 2.55, 2.25, 0.85, "锅", size=15, fill=GOLD,
               color=RGBColor(0x5C, 0x40, 0x20))
        circle(s, left + 1.5, 3.35, 0.6, "A", size=13, fill=TERRA)
        if i == 0:
            line(s, left + 2.1, 3.5, 0.8, color=TERRA, h=0.04)
            text(s, left + 1.3, 2.9, 2.6, 0.3, "等锅", size=13, color=TERRA_D, bold=True)
        else:
            line(s, left + 1.85, 3.4, 0.9, color=SAGE, h=0.04)
            circle(s, left + 0.75, 3.5, 0.5, "O", size=12, fill=ONION, color=RGBColor(0x5C, 0x40, 0x20))
        text(s, left + 0.4, 4.75, 5.1, 0.5, cap, size=19, bold=True, align=PP_ALIGN.CENTER,
             color=TERRA_D if i == 0 else RGBColor(0x4E, 0x6B, 0x5B))
    footer(s, "画面一样、第一步动作也一样——光看现场猜不出来，只能靠印象补那半句。")


def pills(slide, x, y, items, width=0.62, height=0.62, dark=False, gap=0.16):
    """Draw a horizontal row of labels; ``items`` = (kind, text)."""
    cx = x
    for kind, label in items:
        span = max(0.9, 0.19 * len(label) + 0.36)
        if kind == "arrow":
            text(slide, cx, y + 0.12, 0.4, 0.35, label, size=18, color=RGBColor(0xB7, 0xA7, 0x9B), bold=True)
            cx += 0.5
            continue
        if kind == "chip-man":
            box(slide, cx, y + 0.06, span, height - 0.12, label, size=12, bold=True, color=WHITE,
                fill=TERRA, align=PP_ALIGN.CENTER)
        elif kind == "chip-real":
            box(slide, cx, y + 0.06, span, height - 0.12, label, size=12, bold=True, color=WHITE,
                fill=SAGE, align=PP_ALIGN.CENTER)
        elif kind == "acc":
            box(slide, cx, y, span, height, label, size=13, bold=True, color=WHITE,
                fill=TERRA if dark else SAGE, align=PP_ALIGN.CENTER)
        else:
            box(slide, cx, y, span, height, label, size=13, bold=True,
                color=RGBColor(0xEA, 0xDF, 0xD6) if dark else RGBColor(0x6B, 0x4A, 0x3F),
                fill=DARK_CARD if dark else RGBColor(0xEF, 0xE6, 0xDA), align=PP_ALIGN.CENTER)
        cx += span + gap
    return cx


def slide_event_build(prs):
    s = blank(prs)
    bg(s)
    header(s, "造题 · 第一步", "50 个局面，是“摆”出来再“筛”出来的")
    nodes = [("搭布景", "两人站位 + 一口在煮的锅", TERRA),
             ("挨个试", "每种站位都试一遍", SAGE),
             ("两个 Alice 各决定", "想法不同才要", GOLD),
             ("一道筛子", "只留“看不出差别”的", TERRA)]
    x = 0.72
    for i, (title, sub, col) in enumerate(nodes):
        box(s, x, 1.65, 2.45, 1.5, fill=CARD, line=CARD_LINE)
        circle(s, x + 0.24, 1.8, 0.5, str(i + 1), size=15, fill=col)
        text(s, x + 0.24, 2.42, 2.0, 0.35, title, size=14, bold=True)
        text(s, x + 0.24, 2.74, 2.05, 0.35, sub, size=10, color=INK_SOFT)
        if i < 3:
            text(s, x + 2.5, 2.1, 0.4, 0.35, "→", size=18, color=RGBColor(0xC0, 0xAF, 0xA4), bold=True)
        x += 2.87
    # 迷你布景
    box(s, 0.72, 3.35, 5.4, 2.85, fill=CARD, line=CARD_LINE)
    cell = 0.34
    bx, by = 0.95, 3.62
    for gx in range(12):
        line(s, bx + gx * cell, by, 0.008, color=RGBColor(0xEF, 0xE6, 0xDA), h=7 * cell, dash=False)
    for gy in range(8):
        line(s, bx, by + gy * cell, 11 * cell, color=RGBColor(0xEF, 0xE6, 0xDA), dash=False)
    circle(s, bx + 5 * cell, by + 1.4 * cell, 0.5, "锅", size=11, fill=GOLD, color=RGBColor(0x5C, 0x40, 0x20))
    circle(s, bx + 3 * cell, by + 3 * cell, 0.36, "A", size=10, fill=TERRA)
    circle(s, bx + 7.5 * cell, by + 4.4 * cell, 0.36, "B", size=10, fill=SAGE)
    line(s, bx + 6.9 * cell, by + 4.3 * cell, 1.3, color=SAGE)
    text(s, 0.95, 5.82, 5.0, 0.3, "布景：两人空手 · 锅里有 3 个洋葱、正煮着", size=11, color=RGBColor(0x8C, 0x7C, 0x72))
    # 漏斗
    for i, (w, lab, sub, col) in enumerate([
            (5.9, "5,043 种站位", "每张图都试一遍", RGBColor(0xC6, 0xB3, 0xA8)),
            (4.4, "131–182 个合格", "想法不同 + 画面一致", SAGE),
            (3.8, "110 个够用", "分成三份、互不重叠", TERRA)]):
        y = 3.35 + i * 0.85
        box(s, 6.55, y, w, 0.55, lab, size=13, bold=True, color=WHITE, fill=col)
        text(s, 6.55, y + 0.57, 6.1, 0.26, sub, size=10, color=INK_SOFT)
    splits = [("开发 20", "用来调提示词"), ("印象 40", "Bob 旁观、攒印象"), ("正式 50", "正式出题、定结论")]
    sx = 6.55
    for title, sub in splits:
        box(s, sx, 5.95, 1.8, 0.85, fill=RGBColor(0xEF, 0xE6, 0xDA))
        text(s, sx + 0.1, 6.02, 1.6, 0.28, title, size=12, bold=True, align=PP_ALIGN.CENTER)
        text(s, sx + 0.06, 6.32, 1.68, 0.4, sub, size=9, color=INK_SOFT, align=PP_ALIGN.CENTER)
        sx += 1.92
    text(s, 0.95, 6.35, 5.4, 0.6, "筛的规矩：第一步动作相同、目标设施不同、走完这步后 Bob 看到的画面一字不差。",
         size=11, color=INK_SOFT)


def slide_event_run(prs):
    s = blank(prs)
    bg(s)
    header(s, "造题 · 第二步", "一题两段：上半段人造，下半段真跑", accent=SAGE)
    box(s, 0.72, 1.95, 11.9, 1.1, fill=RGBColor(0xFB, 0xED, 0xE9), line=TERRA)
    pills(s, 1.05, 2.19, [("chip-man", "人造"), ("plain", "画面"), ("plain", "第一步动作"),
                          ("plain", "印象"), ("arrow", "→"), ("acc", "Qwen 猜：等锅 / 拿料")])
    text(s, 0.72, 3.25, 11.9, 0.4, "↓ 同一起点 ↓", size=18, bold=True,
         color=RGBColor(0xB7, 0xA7, 0x9B), align=PP_ALIGN.CENTER)
    box(s, 0.72, 3.85, 11.9, 1.1, fill=DARK_CARD, line=RGBColor(0x4A, 0x3E, 0x38))
    pills(s, 1.05, 4.09, [("chip-real", "真实"), ("plain", "Alice（成长后）"),
                          ("plain", "Bob（规则控制器）"), ("arrow", "→"), ("plain", "跑 200 步"),
                          ("arrow", "→"), ("acc", "数出几碗汤")], dark=True)
    text(s, 0.72, 5.3, 11.9, 0.4, "造的那一帧只负责“出题”；答案对不对、配合好不好，全看下面这段真跑。",
         size=13, color=INK_SOFT)


def slide_impression_build(prs):
    s = blank(prs)
    bg(s)
    header(s, "印象 · 第一步", "Bob 的印象不是听说的，是他亲眼看见的")
    nodes = [("40 个起点局面", "另一批，不碰正式题", TERRA),
             ("让 Alice 真跑一段", "最多 20 步", SAGE),
             ("只看 3×3 里有的", "走出视野就断片", GOLD),
             ("攒成两本小本子", "成长前跑 / 成长后跑", TERRA)]
    x = 0.72
    for i, (title, sub, col) in enumerate(nodes):
        box(s, x, 1.65, 2.45, 1.5, fill=CARD, line=CARD_LINE)
        circle(s, x + 0.24, 1.8, 0.5, str(i + 1), size=15, fill=col)
        text(s, x + 0.24, 2.42, 2.05, 0.35, title, size=14, bold=True)
        text(s, x + 0.24, 2.74, 2.05, 0.35, sub, size=10, color=INK_SOFT)
        if i < 3:
            text(s, x + 2.5, 2.1, 0.4, 0.35, "→", size=18, color=RGBColor(0xC0, 0xAF, 0xA4), bold=True)
        x += 2.87
    for i, (title, num, frac, col, cap) in enumerate([
            ("旧印象 · 由成长前跑出来", "27 / 40", 0.675, RGBColor(0xA9, 0x8D, 0x80),
             "片段收场是“一直没去拿洋葱，空手站着”。"),
            ("新印象 · 由成长后跑出来", "31 / 40", 0.775, SAGE,
             "片段里出现了“被看到站在洋葱台旁”。")]):
        left = 0.72 + i * 6.2
        box(s, left, 3.55, 5.9, 1.9, fill=CARD, line=CARD_LINE)
        text(s, left + 0.24, 3.72, 5.4, 0.35, title, size=15, bold=True,
             color=TERRA_D if i == 0 else RGBColor(0x4E, 0x6B, 0x5B))
        text(s, left + 0.24, 4.12, 1.3, 0.35, num, size=14, bold=True, color=RGBColor(0x6B, 0x4A, 0x3F))
        box(s, left + 1.6, 4.1, 4.05, 0.44, fill=RGBColor(0xEF, 0xE6, 0xDA))
        box(s, left + 1.6, 4.1, 4.05 * frac, 0.44, f"{frac:.1%}", size=13, bold=True, color=WHITE,
            fill=col, align=PP_ALIGN.RIGHT)
        text(s, left + 0.24, 4.75, 5.4, 0.5, cap, size=11, color=INK_SOFT)
    text(s, 0.72, 5.75, 11.9, 0.4, "两边用的是同一批 40 个起点（举例：group_0）。看到的东西不一样，两本本子就不一样。",
         size=13, color=INK_SOFT)


def slide_impression_frame(prs):
    s = blank(prs)
    bg(s)
    header(s, "印象 · 结构", "印象里的一帧，只装他看见的东西", accent=SAGE)
    box(s, 0.72, 1.8, 5.9, 3.5, fill=CARD, line=CARD_LINE)
    text(s, 0.98, 1.96, 5.4, 0.35, "他一眼看到的 3 × 3", size=16, bold=True)
    tiles = [("盘台", False), ("墙 / 台面", False), ("锅[空]", False),
             ("伙伴（空手）", True), ("空地", False), ("空地", False),
             ("空地", False), ("空地", False), ("空地", False)]
    cw, ch = 1.72, 0.62
    for idx, (label, hot) in enumerate(tiles):
        cx = 0.98 + (idx % 3) * (cw + 0.08)
        cy = 2.45 + (idx // 3) * (ch + 0.08)
        box(s, cx, cy, cw, ch, label, size=11, bold=hot,
            color=RGBColor(0x3F, 0x5C, 0x4C) if hot else RGBColor(0x5C, 0x4A, 0x42),
            fill=RGBColor(0xDC, 0xE8, 0xE1) if hot else RGBColor(0xFB, 0xF7, 0xF1),
            line=RGBColor(0xBB, 0xD3, 0xC7) if hot else RGBColor(0xE8, 0xDC, 0xCB),
            align=PP_ALIGN.CENTER)
    text(s, 0.98, 4.7, 5.4, 0.5, "绿格就是 Alice 当时站的位置（这一帧她在 Bob 左手边）。", size=11, color=INK_SOFT)
    box(s, 6.9, 1.8, 5.7, 3.5, fill=CARD, line=CARD_LINE)
    text(s, 7.16, 1.96, 5.2, 0.35, "同一帧记下的字段", size=16, bold=True)
    fields = ["Alice 位置 (4, 1)", "朝向 左", "手持 空", "第 1 步",
              "动作 向左移动", "Bob 位置 (5, 1)", "看过的锅 锅[空]"]
    fx, fy = 7.16, 2.45
    for f in fields:
        w = 0.16 * len(f) + 0.4
        if fx + w > 12.4:
            fx, fy = 7.16, fy + 0.62
        box(s, fx, fy, w, 0.5, f, size=11, bold=True, color=RGBColor(0x4A, 0x3B, 0x33),
            fill=RGBColor(0xEF, 0xE6, 0xDA), align=PP_ALIGN.CENTER)
        fx += w + 0.16
    text(s, 7.16, 4.7, 5.2, 0.5, "每条都带“亲眼看到”的来源标记；看不到的东西不会写进来。", size=11, color=INK_SOFT)
    for i, bad in enumerate(["✕ 她的真实意图", "✕ 视野外的位置", "✕ 没看见的锅态"]):
        box(s, 0.72 + i * 4.1, 5.6, 3.85, 0.7, bad, size=13, bold=True, color=TERRA_D,
            fill=RGBColor(0xFB, 0xED, 0xE9), line=TERRA)


def slide_bob(prs):
    s = blank(prs)
    bg(s)
    header(s, "Bob 的配置", "他怎么“看”、怎么“猜”", accent=SAGE)
    cards = [("3 × 3 视野", "只记亲眼看到的：位置、手持、动作。", TERRA, TERRA_L),
             ("Qwen3.5-9B", "温度 0，固定 JSON：意图 / 目标设施 / 置信度。", SAGE, SAGE_L),
             ("每题 3 票", "同一问题问 3 次，取多数，压住抖动。", GOLD, GOLD_L),
             ("规则控制器", "猜完按固定规则干活：供洋葱或做服务。", TERRA, TERRA_L)]
    x = 0.72
    for title, sub, col, colL in cards:
        box(s, x, 1.95, 2.86, 3.6, fill=CARD, line=CARD_LINE)
        circle(s, x + 0.3, 2.2, 0.72, "", fill=colL)
        circle(s, x + 0.46, 2.36, 0.4, "", fill=col)
        text(s, x + 0.3, 3.2, 2.3, 0.5, title, size=17, bold=True)
        text(s, x + 0.3, 3.8, 2.35, 1.5, sub, size=12, color=INK_SOFT)
        x += 3.0
    footer(s, "印象片段最长 20 步；Alice 一走出视野，这段观察就结束。")


def slide_results(prs):
    s = blank(prs)
    bg(s, dark=True)
    header(s, "跑完的结果", "印象一旧，判断就跟着旧", dark=True)
    bars = [("旧印象", 0.2178, "21.8%", RGBColor(0x8E, 0x6A, 0x5E)),
            ("新印象", 0.90, "90.0%", TERRA)]
    for i, (lab, frac, val, col) in enumerate(bars):
        y = 2.15 + i * 1.15
        text(s, 0.72, y + 0.18, 1.2, 0.4, lab, size=16, bold=True, color=RGBColor(0xF3, 0xED, 0xE6))
        box(s, 2.0, y, 10.6, 0.75, fill=DARK_CARD)
        box(s, 2.0, y, max(1.2, 10.6 * frac), 0.75, val, size=16, bold=True, color=WHITE,
            fill=col, align=PP_ALIGN.RIGHT)
    text(s, 0.72, 4.45, 11.9, 0.4, "猜对意图的比例（9 组共 450 个局面，9/9 组方向一致）",
         size=14, color=RGBColor(0xD9, 0xCC, 0xC2))
    box(s, 0.72, 5.0, 5.6, 1.3, fill=DARK_CARD)
    text(s, 1.0, 5.18, 5.0, 0.35, "印象与真身匹配", size=14, color=RGBColor(0xF3, 0xED, 0xE6))
    text(s, 1.0, 5.5, 5.0, 0.6, "84.4%", size=32, bold=True, color=RGBColor(0xF3, 0xED, 0xE6))
    box(s, 6.98, 5.0, 5.6, 1.3, fill=DARK_CARD)
    text(s, 7.26, 5.18, 5.0, 0.35, "印象与真身不匹配", size=14, color=RGBColor(0xF3, 0xED, 0xE6))
    text(s, 7.26, 5.5, 5.0, 0.6, "16.2%", size=32, bold=True, color=RGBColor(0xE0, 0x8B, 0x7E))
    footer(s, "得分只作辅助：新印象比旧印象多 0.39 碗 / +7.9 分。", dark=True)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for fn in (slide_title, slide_board, slide_roles, slide_training,
               slide_training_data, slide_training_why, slide_group_matrix,
               slide_pipeline, slide_key_moment, slide_event_build, slide_event_run,
               slide_impression_build, slide_impression_frame,
               slide_bob, slide_results):
        fn(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"written: {OUT} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
